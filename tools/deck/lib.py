# -*- coding: utf-8 -*-
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY  = C(0x00, 0x2C, 0x5F)
BLUE  = C(0x1F, 0x6F, 0xB2)
SKY   = C(0x90, 0xD4, 0xFF)
TINT  = C(0xEA, 0xF4, 0xFB)
PAPER = C(0xF5, 0xF7, 0xFA)
INK   = C(0x1A, 0x1A, 0x1A)
GRAY  = C(0x6B, 0x72, 0x80)
LINE  = C(0xC9, 0xD6, 0xE3)
RED   = C(0x9B, 0x3B, 0x3B)
GOLD  = C(0xC8, 0x8A, 0x2A)
WHITE = C(0xFF, 0xFF, 0xFF)

KO = "Noto Sans KR"
L, R = 1.458, 25.21
W = R - L
TOP = 2.95


def drop(sh):
    sh._element.getparent().remove(sh._element)


def drop_pics(slide):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            drop(sh)


def set_title(slide, text, width=None):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() in (
                "프로젝트 개관", "프로젝트 목적"):
            if width is not None:
                sh.width = I(width)
            p = sh.text_frame.paragraphs[0]
            for i, r in enumerate(p.runs):
                r.text = text if i == 0 else ""
            return
    raise SystemExit("title not found")


def tb(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return box, tf


def line(tf, text, size, color=INK, bold=False, space_before=0,
         space_after=0, first=False, spacing=1.35, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = spacing
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = KO
    return p


def mixed(tf, parts, size, spacing=1.3, first=True, space_before=0, align=None):
    """parts = [(text, color, bold), ...] 한 문단 안에서 색을 섞는다."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = spacing
    p.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = KO
    return p


def card(slide, x, y, w, h, fill=TINT, border=LINE, radius=0.045):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                I(x), I(y), I(w), I(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def rule(slide, x, y, w, color=LINE, weight=1.2):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w),
                                Pt(weight))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def chip(slide, x, y, d, label, fill=NAVY, fg=WHITE, size=20):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = fg
    r.font.name = KO
    return sh


def arrow(slide, x, y, w, h, color=SKY):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def banner(slide, y, h, parts, size=19):
    card(slide, L, y, W, h, fill=NAVY, border=None)
    _, t = tb(slide, L + 0.5, y, W - 1.0, h, anchor=MSO_ANCHOR.MIDDLE)
    mixed(t, parts, size)


def set_face(run, latin, ea=None):
    """런의 latin / eastAsia 글꼴을 함께 지정한다."""
    from pptx.oxml.ns import qn
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", latin), ("a:ea", ea or latin)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def down_arrow(slide, x, y, w, h, color=SKY):
    sh = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, I(x), I(y), I(w), I(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh
