# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lib import *

prs = Presentation("base.pptx")
S = prs.slides
BOT = 14.1                       # 본문이 넘지 말아야 할 아래 한계

# ═══════════════ 6. 표지 — Keywords 교체 ═══════════════
KW = {"추천시스템": "조기 탐지", "예측모델": "예측 모형",
      "클러스터링": "LLM 활용", "선형계획법": "추천 시스템"}
for sh in S[5].shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text.strip()
        if t in KW:
            runs = sh.text_frame.paragraphs[0].runs
            for i, r in enumerate(runs):
                r.text = KW[t] if i == 0 else ""

# ═══════════════ 8. 머리말의 현대차 흔적 제거 ═══════════════
for sh in S[7].shapes:
    if sh.has_text_frame and "블루멤버스" in sh.text_frame.text:
        p = sh.text_frame.paragraphs[0]
        for i, r in enumerate(p.runs):
            r.text = ("LLM 활용 패션 트렌드 조기 탐지 및 사입 후보 추천 시스템"
                      if i == 0 else "")
            set_face(r, "Inter")

